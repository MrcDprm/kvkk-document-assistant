import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from app import extract
from app.extract import ExtractError, extract_text, normalize_text

MB = 1024 * 1024


def run(filename, data, max_bytes=10 * MB, max_chars=200_000):
    return extract_text(filename, data, max_bytes, max_chars)


def error_code(filename, data, **kwargs):
    with pytest.raises(ExtractError) as info:
        run(filename, data, **kwargs)
    return info.value.code


def make_pdf(*pages: str) -> bytes:
    """A minimal valid PDF with one line of ASCII text per page (Helvetica)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        content_id = len(objects)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs, table=None) -> bytes:
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --- txt ---

def test_txt_utf8():
    assert run("not.txt", "Ayşe Yılmaz, İstanbul".encode("utf-8")) == "Ayşe Yılmaz, İstanbul"


def test_txt_utf8_with_bom():
    assert run("not.txt", "﻿Şule".encode("utf-8")) == "Şule"


def test_txt_turkish_windows_encoding():
    assert run("eski.txt", "Ağaç ışık".encode("cp1254")) == "Ağaç ışık"


def test_txt_binary_is_rejected():
    assert error_code("resim.txt", b"\x89PNG\r\n\x1a\n\x00\x00\x00") == "unsupported_type"


# --- pdf ---

def test_pdf_text_from_all_pages():
    text = run("belge.pdf", make_pdf("Ayse Yilmaz 10000000146", "Ikinci sayfa"))
    assert "Ayse Yilmaz 10000000146" in text
    assert "Ikinci sayfa" in text


def test_pdf_with_password_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="gizli")
    out = io.BytesIO()
    writer.write(out)
    assert error_code("sifreli.pdf", out.getvalue()) == "encrypted"


def test_pdf_without_text_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)  # like a scanned page: no text layer
    out = io.BytesIO()
    writer.write(out)
    assert error_code("tarama.pdf", out.getvalue()) == "no_text"


def test_broken_pdf_is_rejected():
    assert error_code("bozuk.pdf", b"%PDF-1.4\nthis is not a pdf") == "corrupt"


def test_too_many_pages(monkeypatch):
    monkeypatch.setattr(extract, "MAX_PDF_PAGES", 2)
    assert error_code("uzun.pdf", make_pdf("a", "b", "c")) == "too_many_pages"


# --- docx ---

def test_docx_paragraphs_and_table_in_order():
    data = make_docx(["Kira sözleşmesi", "Taraflar:"], table=[["Ad", "Telefon"], ["Ayşe Yılmaz", "0532 123 45 67"]])
    assert run("sozlesme.docx", data) == (
        "Kira sözleşmesi\nTaraflar:\nAd | Telefon\nAyşe Yılmaz | 0532 123 45 67"
    )


def test_zip_that_is_not_docx_is_rejected():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hello")
    assert error_code("arsiv.docx", out.getvalue()) == "unsupported_type"


def test_zip_bomb_is_rejected(monkeypatch):
    monkeypatch.setattr(extract, "MAX_DOCX_UNCOMPRESSED_BYTES", 1 * MB)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (5 * MB))
    data = out.getvalue()
    assert len(data) < 50_000  # tiny on disk...
    assert error_code("bomba.docx", data) == "too_large"  # ...but refused before unpacking


def test_broken_docx_is_rejected():
    assert error_code("bozuk.docx", b"PK\x03\x04garbage") == "corrupt"


# --- type and size checks ---

@pytest.mark.parametrize("filename", ["resim.png", "betik.exe", "dosya", "tablo.xlsx"])
def test_unsupported_extension(filename):
    assert error_code(filename, b"hello") == "unsupported_type"


def test_extension_must_match_content():
    assert error_code("belge.txt", make_pdf("x")) == "unsupported_type"
    assert error_code("belge.pdf", b"plain text") == "unsupported_type"


def test_extension_is_case_insensitive():
    assert run("NOT.TXT", b"merhaba") == "merhaba"


def test_file_too_large():
    assert error_code("buyuk.txt", b"a" * 101, max_bytes=100) == "too_large"


def test_text_too_long():
    assert error_code("uzun.txt", b"a" * 101, max_chars=100) == "text_too_long"


# --- normalize_text (also used for pasted text) ---

def test_normalize_joins_combining_characters():
    decomposed = "ş"  # "s" + combining cedilla, as some PDFs store "ş"
    assert normalize_text(f"Ay{decomposed}e", 100) == "Ayşe"


def test_normalize_line_endings_and_blank_lines():
    assert normalize_text("a  \r\nb\r\n\r\n\r\n\r\nc\x00", 100) == "a\nb\n\nc"


def test_normalize_empty_text():
    with pytest.raises(ExtractError) as info:
        normalize_text("   \n\n ", 100)
    assert info.value.code == "no_text"
